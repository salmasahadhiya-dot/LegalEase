import os, io, requests, streamlit as st
from datetime import date
from docx import Document
from fpdf import FPDF
from dotenv import load_dotenv

load_dotenv()
st.set_page_config(page_title="LegalEase", page_icon="⚖️", layout="centered")
st.title("⚖️ LegalEase")
st.subheader("AI-Powered Legal Document Generator")
st.caption("AI-generated draft only. Please have a qualified lawyer review important legal documents.")
backend_url = os.getenv("BACKEND_URL", "http://127.0.0.1:8000").rstrip("/")

with st.form("document_form"):
    document_type = st.selectbox("Document type", ["Employment Contract", "Non-Disclosure Agreement (NDA)", "Lease Agreement", "Freelance Work Contract", "Employment Offer Letter", "Service Agreement", "General Agreement"])
    parties = st.text_area("Parties involved", placeholder="Example: Jane (Freelancer), ABC Ltd (Client)")
    terms = st.text_area("Important clauses / terms", placeholder="Separate each term with a semicolon (;)")
    effective_date = st.date_input("Effective date", value=date.today())
    submitted = st.form_submit_button("Generate Document")

if submitted:
    if not parties.strip() or not terms.strip():
        st.error("Please enter the parties involved and the terms.")
    else:
        try:
            payload = {"document_type": document_type, "parties": parties.strip(), "terms": terms.strip(), "effective_date": effective_date.isoformat()}
            with st.spinner("Generating your draft..."):
                response = requests.post(f"{backend_url}/generate", json=payload, timeout=120)
                response.raise_for_status()
                st.session_state["generated_document"] = response.json()["document"]
                st.session_state["generated_type"] = document_type
        except requests.RequestException as exc:
            st.error("Unable to contact backend. Check that the API service is running and BACKEND_URL is correct.")
            st.caption(str(exc))
        except (KeyError, ValueError):
            st.error("Backend returned an unexpected response. Check service logs.")

if "generated_document" in st.session_state:
    st.markdown("### Document preview")
    edited_text = st.text_area("Edit the generated document before downloading", value=st.session_state["generated_document"], height=380, key="edited_document")
    st.warning("This is an AI-generated draft, not legal advice. Verify it with a qualified lawyer before use.")
    st.download_button("Download TXT", data=edited_text.encode("utf-8"), file_name="legalease_document.txt", mime="text/plain")
    doc = Document()
    doc.add_heading(st.session_state.get("generated_type", "Legal Document"), 0)
    for paragraph in edited_text.splitlines():
        doc.add_paragraph(paragraph)
    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    st.download_button("Download DOCX", data=docx_buffer.getvalue(), file_name="legalease_document.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in edited_text.splitlines():
        pdf.multi_cell(0, 7, line.encode("latin-1", "replace").decode("latin-1"))
    st.download_button("Download PDF", data=bytes(pdf.output()), file_name="legalease_document.pdf", mime="application/pdf")
